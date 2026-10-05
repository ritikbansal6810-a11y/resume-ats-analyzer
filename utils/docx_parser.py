"""
docx_parser.py
──────────────
Extract plain text from DOCX files using python-docx.
Handles:
  - Regular DOCX paragraphs and tables
  - Headers / footers (best-effort)
  - Empty or corrupted DOCX files
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Optional

try:
    from docx import Document
    from docx.oxml.ns import qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB


@dataclass
class DOCXParseResult:
    text: str = ""
    warnings: list[str] = field(default_factory=list)
    error: Optional[str] = None

    @property
    def success(self) -> bool:
        return self.error is None and bool(self.text.strip())


def _para_text(para) -> str:
    """Return the text of a paragraph, including runs inside hyperlinks."""
    parts: list[str] = []
    for elem in para._p.iter():
        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
        if tag == "t" and elem.text:
            parts.append(elem.text)
        elif tag == "br":
            parts.append("\n")
    return "".join(parts)


def _extract_table_text(table) -> str:
    """Render a DOCX table as plain text rows."""
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_text_from_docx(file_bytes: bytes) -> DOCXParseResult:
    """
    Extract readable text from a DOCX file supplied as raw bytes.

    Returns a DOCXParseResult with:
        .text     – extracted plain text
        .warnings – non-fatal issues
        .error    – fatal error message (None on success)
    """
    result = DOCXParseResult()

    # ── Dependency check ───────────────────────────────────────────────────
    if not DOCX_AVAILABLE:
        result.error = (
            "python-docx is not installed. Run: pip install python-docx"
        )
        return result

    # ── Size / empty guard ─────────────────────────────────────────────────
    if len(file_bytes) == 0:
        result.error = "The uploaded DOCX file is empty (0 bytes)."
        return result

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        result.error = (
            f"File is too large ({len(file_bytes) / 1024 / 1024:.1f} MB). "
            f"Maximum allowed size is {MAX_FILE_SIZE_BYTES // 1024 // 1024} MB."
        )
        return result

    # ── Open document ──────────────────────────────────────────────────────
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        result.error = (
            f"Could not open DOCX file: {exc}. "
            "The file may be corrupted or in an unsupported format."
        )
        return result

    parts: list[str] = []

    # ── Body paragraphs ────────────────────────────────────────────────────
    for para in doc.paragraphs:
        text = _para_text(para).strip()
        if text:
            parts.append(text)

    # ── Tables ─────────────────────────────────────────────────────────────
    if doc.tables:
        parts.append("\n[TABLE CONTENT]")
        for table in doc.tables:
            parts.append(_extract_table_text(table))

    # ── Headers / Footers ─────────────────────────────────────────────────
    header_footer_texts: list[str] = []
    for section in doc.sections:
        for hf in [section.header, section.footer]:
            if hf is not None:
                for para in hf.paragraphs:
                    t = para.text.strip()
                    if t:
                        header_footer_texts.append(t)

    if header_footer_texts:
        parts.append("\n[HEADER/FOOTER CONTENT]")
        parts.extend(header_footer_texts)
        result.warnings.append(
            "This document contains headers/footers. "
            "Their content has been extracted but some ATS systems may ignore them."
        )

    result.text = "\n".join(parts)

    if not result.text.strip():
        result.error = "No readable text could be extracted from the DOCX file."

    return result
